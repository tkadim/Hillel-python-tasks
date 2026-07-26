import openpyxl

STORE_DATA_FILE = "store_data.xlsx"

class Product:
    def __init__(self, title: str, category: str, price: float, quantity: int):
        self.title = title
        self.category = category
        self.price = price
        self.quantity = quantity

    def change_price(self, new_price: float):
        self.price = new_price

    def change_stock_quantity(self, new_quantity: int):
        self.quantity = new_quantity

    def __str__(self):
        return f"{self.title}"


class Customer:
    def __init__(self, name: str, email: str):
        self.name = name
        self.email = email
        self.order_list: list[Order] = []

    def add_order(self, new_order: Order):
        self.order_list.append(new_order)

    def __str__(self):
        return f"Customer Name: {self.name}\nOrders List:\n{[order.__str__() for order in self.order_list]}"


class Order:
    def __init__(self):
        self.products_list:list[Product] = []
        self.total_sum = 0

    def add_product(self, product: Product, quantity: int = 1):
        if product.quantity >= quantity:
            for _ in range(quantity):
                self.products_list.append(product)
            product.quantity -= quantity
        else:
            raise ValueError(f"Product {product.title} is out of stock!")

    def calculate_total(self):
        self.total_sum = sum(prod.price for prod in self.products_list)
        return self.total_sum

    def __str__(self):
        return f"Products List: {[prod.__str__() for prod in self.products_list]} Total Sum: {self.total_sum}"


class Store:
    def __init__(self):
        self.products: list[Product] = []
        self.customers: list[Customer] = []
        self.orders = {}

    def add_customer(self, name: str, email: str):
        customer = Customer(name, email)
        self.customers.append(customer)
        return customer

    def add_product(self, title: str, category: str, price: float, quantity: int):
        product = Product(title, category, price, quantity)
        self.products.append(product)
        return product

    def add_order(self, customer: Customer, products_list: list[Product]):
        order = Order()
        for prod in products_list:
            try:
                order.add_product(prod)
            except ValueError as e:
                print(f"[ERROR]: {e}")

        if order.products_list:
            order.calculate_total()

            if customer.email not in list(self.orders.keys()):
                self.orders[customer.email] = [order]
            else:
                self.orders[customer.email].append(order)

    def get_orders_for(self, customer: Customer):
        return self.orders[customer.email]

    def load_data_from_file(self, file_name: str):
        wb = openpyxl.load_workbook(file_name)
        ws = wb["Products"]

        for col in ws.iter_rows(values_only=True):
            if not any(col):
                continue

            self.add_product(str(col[1]), str(col[2]), float(col[3]), int(col[4]))

        ws = wb["Customers"]

        for col in ws.iter_rows(values_only=True):
            if not any(col):
                continue

            self.add_customer(str(col[1]), str(col[2]))


store = Store()
store.load_data_from_file(STORE_DATA_FILE)

customer1 = store.customers[0]
customer2 = store.add_customer('Tom', "tom@gmail.com")

print([customer.name for customer in store.customers])
print("---------------------------------")

prod1 = store.products[3]
prod2 = store.add_product("Mouse", "Electronics", 20, 1)

print(f"Stock quantity of product \"{prod1.title}\": {prod1.quantity}")
print("---------------------------------")

store.add_order(customer1, [prod1, prod2])

print(f"Stock quantity of product \"{prod1.title}\": {prod1.quantity}")
print("---------------------------------")

store.add_order(customer1, [prod1])

print(f"Stock quantity of product \"{prod1.title}\": {prod1.quantity}")
print("---------------------------------")

print(f"Store Orders Info for Customer1: {[ord.__str__() for ord in store.get_orders_for(customer1)]}")
print("---------------------------------")

store.add_order(customer1, [prod1])

print(f"Store Orders Info for Customer1: {[ord.__str__() for ord in store.get_orders_for(customer1)]}")
print("---------------------------------")




